"""Deterministic XLSX/CSV importer (spec section 5.5). No model call, ever (AC-6).

Values are taken as printed: Excel floats go through `Decimal(repr(value))` at the boundary and are then
quantised to the cell's number format (`#,##0.00` prints two decimals; `General` prints an integral float
without `.0`). No arithmetic on floats and no computed totals: an absent cell stays an empty string.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from decimal import Decimal
from typing import Literal

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ai.agents.extraction.normalize import normalize_text, normalize_value
from ai.agents.extraction.schema import (
    ExtractedInvoice,
    InvoiceLine,
    TaxSubtotal,
)
from ai.agents.extraction.tabular.template import (
    HEADER_COLUMNS,
    LINE_COLUMNS,
    TAX_COLUMNS,
    match_header,
    row_type,
)

MAX_ROWS = 50_000
MAX_INVOICES = 5_000
MAX_ZIP_ENTRIES = 1_000
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
_HEADER_SCAN_ROWS = 20
_DECIMALS = re.compile(r"\.(0+)")

type TabularFormat = Literal["xlsx", "csv"]


class TabularRejected(Exception):
    """The file cannot be imported at all; `code` is one of the class constants below."""

    ZIP_BOMB = "zip_bomb"
    NOT_A_WORKBOOK = "not_a_workbook"
    TOO_MANY_ROWS = "too_many_rows"
    TOO_MANY_INVOICES = "too_many_invoices"
    EMPTY = "empty"

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass(frozen=True, slots=True)
class ImportedInvoice:
    source_ordinal: int  # 0-based order of first appearance in the file
    source_ref: str  # "rows 2-5" (1-based sheet/file rows)
    invoice: ExtractedInvoice  # normalised


@dataclass(frozen=True, slots=True)
class ImportResult:
    method: TabularFormat
    invoices: tuple[ImportedInvoice, ...]
    missing_columns: tuple[str, ...] = ()  # non-empty -> needs_review/import_mapping_incomplete, no invoices


# ------------------------------------------------------------------ cell values
def _places(number_format: str) -> int | None:
    """Decimal places a number format prints ('#,##0.00' -> 2, '0' -> 0); None for General/unknown."""
    fmt = number_format.split(";")[0]
    if fmt in ("", "General", "@"):
        return None
    if m := _DECIMALS.search(fmt):
        return len(m[1])
    return 0 if "0" in fmt else None


def cell_text(value: object, number_format: str = "General") -> str:
    """One cell as the text it prints. Floats never reach arithmetic: Decimal(repr(value)) first."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, int | float):  # openpyxl reads a whole-number float back as int
        d = Decimal(repr(value)) if isinstance(value, float) else Decimal(value)
        places = _places(number_format)
        if places is not None:
            d = d.quantize(Decimal(1).scaleb(-places))
        elif d == d.to_integral_value():
            d = d.quantize(Decimal(1))
        return format(d, "f")
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value)


# ------------------------------------------------------------------ readers
def check_zip(data: bytes) -> None:
    """Zip-bomb guard before openpyxl touches the file. ZipExtFile never reads past a member's declared
    size and verifies its CRC, so the declared sizes checked here bound what openpyxl can decompress."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile as exc:
        raise TabularRejected(TabularRejected.NOT_A_WORKBOOK) from exc
    if len(infos) > MAX_ZIP_ENTRIES or sum(i.file_size for i in infos) > MAX_UNCOMPRESSED_BYTES:
        raise TabularRejected(TabularRejected.ZIP_BOMB)


def _xlsx_rows(data: bytes) -> Iterator[list[str]]:
    check_zip(data)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError, zipfile.BadZipFile) as exc:
        raise TabularRejected(TabularRejected.NOT_A_WORKBOOK) from exc
    try:
        ws = wb.worksheets[0]
        for row in ws.iter_rows():
            yield [cell_text(getattr(c, "value", None), getattr(c, "number_format", "General") or "General")
                   for c in row]
    finally:
        wb.close()


def decode_csv(data: bytes) -> str:
    """UTF-8 (with or without BOM), else Windows-1256 (Arabic Excel exports)."""
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1256", errors="replace")


def _csv_rows(data: bytes) -> Iterator[list[str]]:
    text = decode_csv(data)
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    yield from csv.reader(io.StringIO(text), dialect)


# ------------------------------------------------------------------ assembly
def _set(obj: object, dotted: str, value: str) -> None:
    *parents, last = dotted.split(".")
    for p in parents:
        obj = getattr(obj, p)
    setattr(obj, last, value)


class _Group:
    def __init__(self, ordinal: int, first_row: int) -> None:
        self.ordinal = ordinal
        self.first_row = first_row
        self.last_row = first_row
        self.invoice = ExtractedInvoice()

    def add(self, rownum: int, cells: dict[str, str]) -> None:
        self.last_row = rownum
        inv = self.invoice
        for col, path in HEADER_COLUMNS.items():
            v = cells.get(col, "")
            if v and not _get(inv, path):  # first non-empty value of the group wins
                _set(inv, path, normalize_value(path, v))
        if row_type(cells.get("row_type", "")) == "tax":
            t = TaxSubtotal()
            for col, field in TAX_COLUMNS.items():
                _set(t, field, normalize_value(field, cells.get(col, "")))
            inv.tax_breakdown.append(t)
        else:
            ln = InvoiceLine()
            for col, field in LINE_COLUMNS.items():
                _set(ln, field, normalize_value(field, cells.get(col, "")))
            inv.lines.append(ln)


def _get(obj: object, dotted: str) -> str:
    for part in dotted.split("."):
        obj = getattr(obj, part)
    return obj if isinstance(obj, str) else ""


def _assemble(method: TabularFormat, rows: Iterable[list[str]]) -> ImportResult:
    it = iter(rows)
    mapping: dict[int, str] | None = None
    missing: tuple[str, ...] = ()
    header_row = 0
    for n, raw in enumerate(it, start=1):
        header_row = n
        cells = [normalize_text(c) for c in raw]
        if sum(1 for c in cells if c) >= 2:
            mapping, missing = match_header(cells)
            break
        if n >= _HEADER_SCAN_ROWS:
            break
    if mapping is None:
        raise TabularRejected(TabularRejected.EMPTY)
    if missing:
        return ImportResult(method=method, invoices=(), missing_columns=missing)

    groups: dict[str, _Group] = {}
    count = 0
    for rownum, raw in enumerate(it, start=header_row + 1):
        cells = {col: normalize_text(raw[i]) for i, col in mapping.items() if i < len(raw)}
        if not any(cells.values()):
            continue
        count += 1
        if count > MAX_ROWS:
            raise TabularRejected(TabularRejected.TOO_MANY_ROWS)
        key = normalize_value("invoice_number", cells.get("invoice_number", ""))
        g = groups.get(key)
        if g is None:
            if len(groups) >= MAX_INVOICES:
                raise TabularRejected(TabularRejected.TOO_MANY_INVOICES)
            g = groups[key] = _Group(len(groups), rownum)
        g.add(rownum, cells)
    if not groups:
        raise TabularRejected(TabularRejected.EMPTY)
    return ImportResult(method=method, invoices=tuple(
        ImportedInvoice(g.ordinal, f"rows {g.first_row}-{g.last_row}", g.invoice) for g in groups.values()))


def import_tabular(data: bytes, fmt: TabularFormat) -> ImportResult:
    """Parses one XLSX or CSV file in the import template. Raises TabularRejected for unreadable files."""
    rows = _xlsx_rows(data) if fmt == "xlsx" else _csv_rows(data)
    return _assemble(fmt, rows)
